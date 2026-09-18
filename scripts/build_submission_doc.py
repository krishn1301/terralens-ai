from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).parents[1]
LIVE_URL = "https://hmkz0x00.github.io/terralens-ai/"
REPO_URL = "https://github.com/Hmkz0x00/terralens-ai"
API_URL = "https://terralens-ai-api.onrender.com"
DOCS_URL = f"{API_URL}/docs"
OUTPUT = ROOT / "submission" / "TerraLens_AI_Submission.docx"
GREEN = "17382D"
PALE = "E7EBDD"
LINE = "D9D9D9"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_margins(cell, value=120):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)


def set_table_borders(table):
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), LINE)
        borders.append(node)
    properties.append(borders)


def keep_row_together(row):
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def style_table(table, widths):
    table.autofit = False
    set_table_borders(table)
    for row_index, row in enumerate(table.rows):
        keep_row_together(row)
        if row_index == 0:
            repeat_header(row)
        for index, cell in enumerate(row.cells):
            cell.width = widths[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, GREEN)
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.bold = True
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE)


def add_hyperlink(paragraph, url):
    relationship = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5C45")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = url
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            if value.startswith("https://"):
                add_hyperlink(cells[index].paragraphs[0], value)
            else:
                cells[index].text = value
    style_table(table, widths)
    # Tables are short: keep each one on a single page so a heading and header
    # row are never stranded at the bottom of a page.
    for row in table.rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    doc.add_paragraph()
    return table


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in (("Title", 28), ("Subtitle", 12), ("Heading 1", 17), ("Heading 2", 12)):
        style = styles[name]
        style.font.name = "Aptos Display" if name != "Normal" else "Aptos"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(size)
        style.font.bold = name.startswith("Heading")
    styles["Title"].paragraph_format.space_after = Pt(8)
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(7)
    styles["Heading 1"].paragraph_format.keep_with_next = True
    styles["Heading 2"].paragraph_format.keep_with_next = True

    title = doc.add_paragraph(style="Title")
    title.add_run("TerraLens AI Project Submission")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Darukaa Earth AI Biodiversity Intelligence Chatbot Challenge")
    intro = doc.add_paragraph()
    intro.add_run("Submission summary. ").bold = True
    intro.add_run(
        "TerraLens AI is a completed, publicly deployed full stack prototype that retrieves indexed "
        "environmental evidence and reasons across soil, water, climate, habitat, biodiversity, and "
        "human pressure. It returns measurable recommendations with explicit sources, confidence, "
        "time horizons, caveats, and a visible reasoning trace."
    )

    doc.add_heading("Submission links and access", level=1)
    add_table(
        doc,
        ["Item", "Link or reviewer instruction"],
        [
            ("Live demo", LIVE_URL),
            ("GitHub repository", REPO_URL),
            ("API documentation", DOCS_URL),
            ("Backend API", API_URL),
            ("Credentials", "None required. No login, account, or API key is needed to use the demo or the API."),
        ],
        [Inches(1.55), Inches(5.1)],
    )
    note = doc.add_paragraph()
    note.add_run("Reviewer note. ").bold = True
    note.add_run(
        "The API runs on Render's free tier, which sleeps after about 15 minutes without traffic. "
        "The first request after a pause can take up to a minute. Opening "
        f"{API_URL}/api/health first wakes it; if a request times out, the interface keeps all inputs "
        "and the reviewer can simply run the assessment again."
    )

    doc.add_heading("Project overview", level=1)
    add_bullets(
        doc,
        [
            "React and TypeScript reviewer workspace with structured landscape inputs and sample cases.",
            "FastAPI service with bounded multi-turn session memory and explicit validation.",
            "SQLite FTS5 retrieval over a structured FAO, IPCC, IPBES, and USDA evidence corpus.",
            "Deterministic multi-metric reasoning that requires at least three environmental variables.",
            "Recommendation output with action, mechanism, impacted metrics, monitoring range, time horizon, confidence, evidence, and caveats.",
            "Docker, automated tests, accessibility checks, responsive browser verification, and GitHub Actions CI.",
        ],
    )

    doc.add_heading("Challenge requirement coverage", level=1)
    add_table(
        doc,
        ["Requirement", "Implementation evidence"],
        [
            ("Knowledge system", "Structured evidence JSON indexed in SQLite FTS5. Retrieval ranks metric, intervention, and condition matches and returns stable source IDs and URLs."),
            ("Conversational intelligence", "The assistant requests missing high-value variables and merges follow-up values into the same conversation session."),
            ("Evidence-backed recommendations", "Every intervention contains at least one resolvable evidence ID and exposes the claim, publisher, year, and URL."),
            ("Multi-metric reasoning", "Rules require at least three supplied variables and model compound constraints such as rainfall, carbon, and monoculture."),
            ("Input handling", "Free text and typed JSON landscape profiles are supported. Missing values stay absent rather than being guessed."),
            ("Output quality", "Responses include recommendations, metric direction, measurable monitoring targets, time horizon, confidence, caveats, and reasoning trace."),
        ],
        [Inches(1.7), Inches(4.95)],
    )

    doc.add_heading("Architecture and data flow", level=1)
    architecture = doc.add_paragraph()
    architecture.add_run("Client. ").bold = True
    architecture.add_run("React collects narrative and structured inputs, preserves drafts on error, and renders clarification or assessment states.")
    architecture = doc.add_paragraph()
    architecture.add_run("API. ").bold = True
    architecture.add_run("FastAPI validates bounds, creates or resumes a session, and merges new profile fields with prior context.")
    architecture = doc.add_paragraph()
    architecture.add_run("Retrieval. ").bold = True
    architecture.add_run("KnowledgeStore expands environmental terms, queries SQLite FTS5, reranks results, and resolves exact evidence records.")
    architecture = doc.add_paragraph()
    architecture.add_run("Reasoning. ").bold = True
    architecture.add_run("ReasoningEngine checks data sufficiency, detects interacting pressures, applies suitability guards, ranks interventions, and constructs a traceable response.")

    # Start the compact schema as a complete block instead of letting Word split
    # its repeated header across a page boundary.
    doc.add_page_break()
    doc.add_heading("Knowledge schema", level=2)
    add_table(
        doc,
        ["Record field", "Purpose"],
        [
            ("id", "Stable reference used by recommendations"),
            ("organization title year url", "Human-verifiable source provenance"),
            ("claim", "Supported scientific mechanism or benchmark"),
            ("metrics", "Environmental outcomes linked to the claim"),
            ("interventions", "Actions supported by the record"),
            ("conditions", "Landscape contexts where the evidence applies"),
        ],
        [Inches(2.0), Inches(4.65)],
    )

    # Let Word paginate naturally here. A forced section break after the
    # preceding table can strand the final schema rows on an otherwise empty page.
    doc.add_heading("Local setup", level=1)
    doc.add_paragraph("Recommended prerequisites are Python 3.12, Node.js 22 or newer, pnpm, and Docker Desktop if using containers.")
    steps = [
        f"Clone {REPO_URL} and open a terminal in the repository folder.",
        "Run docker compose up --build.",
        "Open http://localhost:4173 for the application.",
        "Open http://localhost:8000/docs for interactive API documentation.",
        "Choose Semi-arid wheat farm and select Run grounded assessment.",
    ]
    for text in steps:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.add_run(text)

    doc.add_heading("Development setup", level=2)
    add_table(
        doc,
        ["Area", "Command"],
        [
            ("Python environment", "python -m venv .venv"),
            ("Backend dependencies", ".venv/Scripts/python -m pip install -r backend/requirements-dev.txt"),
            ("Start API", ".venv/Scripts/python -m uvicorn app.main:app --app-dir backend --reload"),
            ("Frontend dependencies", "cd frontend then pnpm install"),
            ("Start client", "pnpm dev"),
        ],
        [Inches(1.65), Inches(5.0)],
    )

    doc.add_heading("Verification evidence", level=1)
    add_table(
        doc,
        ["Quality gate", "Verified result"],
        [
            ("Backend", "26 pytest cases passed; Ruff reported all checks passed."),
            ("Frontend", "7 Vitest cases passed; TypeScript compilation and ESLint completed without errors."),
            ("Production build", "Vite created the optimized production bundle."),
            ("Deployed API", "Health, scenarios, documentation, clarification, merged follow-up assessment, invalid input rejection, session reset, and the CORS allowlist were verified over HTTPS."),
            ("Browser workflow", "On the live demo and locally: sample assessment, clarification then multi-turn follow-up, evidence links, reasoning trace, and recovery from HTTP 503 and an unreachable API passed with no console errors or failed requests."),
            ("Accessibility", "No serious or critical axe violations; keyboard focus remained visible."),
            ("Responsive layout", "No horizontal overflow at 1440 by 900 or 390 by 844."),
            ("Continuous integration", "GitHub Actions CI and the GitHub Pages deployment passed on the final commit."),
        ],
        [Inches(1.6), Inches(5.05)],
    )

    doc.add_heading("CI and deployment", level=1)
    doc.add_paragraph(
        "GitHub Actions runs backend tests and lint, frontend unit tests and lint, and a production build on every push. "
        "A second workflow builds the frontend with the production API URL and publishes it to GitHub Pages. "
        "The backend is a Render Blueprint (render.yaml) for a free Docker web service that runs as a non-root user, "
        "binds to the platform port, and redeploys automatically from the main branch. CORS allows only the GitHub Pages "
        "origin and local development origins, and unexpected server errors return a generic message without internal details. "
        "Docker Compose remains the fastest local reviewer path."
    )

    doc.add_heading("Important limitations", level=1)
    add_bullets(
        doc,
        [
            "The curated corpus is intentionally compact and should be expanded and periodically reviewed for production use.",
            "Impact ranges are planning benchmarks, not site-specific predictions.",
            "Local species selection and implementation timing require regional ecological validation.",
            "Sessions are held in process memory; on the free hosted tier they reset when the service sleeps or redeploys.",
            "The free hosting tier adds a cold start of up to about a minute after inactivity.",
            "Production deployment should add authentication, rate limiting, observability, and managed backups.",
        ],
    )

    doc.add_heading("Reviewer path", level=1)
    doc.add_paragraph(
        f"Open {LIVE_URL}. Run the semi-arid case, inspect the compound recommendations and their metric impacts, expand the reasoning trace, and open an IPCC or FAO source. Then start a new case with only land use to verify that the system asks a targeted clarification instead of producing an unsupported generic answer; adding the requested values continues the same session to a full assessment."
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("TerraLens AI Project Submission  |  Darukaa Earth Challenge")
    footer_run.font.name = "Aptos"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(80, 80, 80)

    doc.core_properties.title = "TerraLens AI Project Submission"
    doc.core_properties.subject = "Darukaa Earth AI Biodiversity Intelligence Chatbot Challenge"
    doc.core_properties.author = "TerraLens AI Candidate"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
